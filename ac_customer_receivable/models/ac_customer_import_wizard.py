import base64
import csv
import io

from odoo import _, api, fields, models
from odoo.exceptions import UserError
from odoo.tools import float_is_zero
from odoo.tools.float_utils import float_round
from odoo.tools.misc import xlsxwriter


# Importable customer master columns (excluding photo / computed / relational).
IMPORT_COLUMNS = [
    ('name', 'Name'),
    ('contact_number', 'Contact Number'),
    ('opening_balance', 'Opening Balance'),
    ('customer_company', 'Customer Company'),
    ('email', 'Email'),
    ('aadhar_no', 'Aadhar Card No'),
    ('pan_no', 'PAN Card No'),
    ('house_no', 'House No/Flat No'),
    ('street', 'Street Name'),
    ('street2', 'Street 2'),
    ('city', 'City'),
    ('state', 'State'),
    ('country', 'Country'),
    ('zip', 'PINCODE'),
    ('bank_account_number', 'Account Number'),
    ('bank_ifsc_code', 'IFSC Code'),
]

MANDATORY_COLUMNS = ('name', 'contact_number', 'opening_balance')

HEADER_ALIASES = {
    'name': ('name', 'customer_name', 'customer'),
    'contact_number': ('contact_number', 'phone', 'mobile', 'contact', 'phone_number'),
    'opening_balance': ('opening_balance', 'balance', 'receivable', 'opening'),
    'customer_company': ('customer_company', 'company', 'company_name'),
    'email': ('email', 'e_mail', 'mail'),
    'aadhar_no': ('aadhar_no', 'aadhaar_no', 'aadhar', 'aadhaar'),
    'pan_no': ('pan_no', 'pan'),
    'house_no': ('house_no', 'house', 'flat_no', 'flat'),
    'street': ('street', 'street_name'),
    'street2': ('street2', 'street_2'),
    'city': ('city',),
    'state': ('state', 'state_name', 'state_code'),
    'country': ('country', 'country_name', 'country_code'),
    'zip': ('zip', 'pincode', 'pin_code', 'postal_code'),
    'bank_account_number': ('bank_account_number', 'account_number', 'account_no'),
    'bank_ifsc_code': ('bank_ifsc_code', 'ifsc_code', 'ifsc'),
}


class AcCustomerImportWizard(models.TransientModel):
    _name = 'ac.customer.import.wizard'
    _description = 'Import Customers with Opening Balance'

    data_file = fields.Binary(string='Excel / CSV File')
    filename = fields.Char(string='Filename')
    template_file = fields.Binary(string='Template File')
    template_filename = fields.Char(string='Template Filename')
    company_id = fields.Many2one(
        'res.company',
        string='Company',
        required=True,
        default=lambda self: self.env.company,
    )
    line_ids = fields.One2many(
        'ac.customer.import.wizard.line',
        'wizard_id',
        string='Preview',
    )
    state = fields.Selection(
        selection=[
            ('upload', 'Upload'),
            ('preview', 'Preview'),
        ],
        default='upload',
        required=True,
    )

    @api.model
    def _template_headers(self):
        return [label for _key, label in IMPORT_COLUMNS]

    @api.model
    def _template_sample_row(self):
        return [
            'John Doe',
            '9876543210',
            '5000',
            'Acme Corp',
            'john@example.com',
            '',
            '',
            '12A',
            'Main Street',
            '',
            'Mumbai',
            'Maharashtra',
            'IN',
            '400001',
            '',
            '',
        ]

    def action_download_template(self):
        self.ensure_one()
        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {'in_memory': True})
        sheet = workbook.add_worksheet('Customers')
        header_style = workbook.add_format({'bold': True, 'bg_color': '#D9E2F3'})
        mandatory_style = workbook.add_format({'bold': True, 'bg_color': '#F8CBAD'})

        headers = self._template_headers()
        sample = self._template_sample_row()
        for col, header in enumerate(headers):
            style = mandatory_style if col < 3 else header_style
            sheet.write(0, col, header, style)
            sheet.write(1, col, sample[col] if col < len(sample) else '')
            sheet.set_column(col, col, max(14, len(header) + 2))

        note = workbook.add_worksheet('Instructions')
        note.write(0, 0, 'Mandatory columns (orange header): Name, Contact Number, Opening Balance')
        note.write(1, 0, 'Positive Opening Balance = receivable (debit). Negative = credit balance.')
        note.write(2, 0, 'Country can be code (IN) or name. State can be name or code for that country.')
        note.write(3, 0, 'Delete the sample row before importing your data (or overwrite it).')
        note.set_column(0, 0, 90)
        workbook.close()

        self.write({
            'template_file': base64.b64encode(output.getvalue()),
            'template_filename': 'customer_import_template.xlsx',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': (
                '/web/content/?model=%s&id=%s&field=template_file'
                '&filename_field=template_filename&download=true'
            ) % (self._name, self.id),
            'target': 'self',
        }

    def action_load_file(self):
        self.ensure_one()
        if not self.data_file:
            raise UserError(_('Please upload an Excel or CSV file.'))

        rows = self._read_import_rows()
        if not rows:
            raise UserError(_('No data rows found in the file.'))

        header_map = self._map_headers(rows[0].keys())
        for key in MANDATORY_COLUMNS:
            if key not in header_map:
                raise UserError(_(
                    'File must include mandatory columns: Name, Contact Number, Opening Balance. '
                    'Download the template for the full format.'
                ))

        self.line_ids.unlink()
        lines_vals = []
        for row_number, row in enumerate(rows, start=2):
            values = self._extract_row_values(row, header_map)
            name = values.get('name') or ''
            phone = values.get('contact_number') or ''
            if not name and not phone and values.get('opening_balance') in (None, ''):
                continue

            missing = []
            if not name:
                missing.append('Name')
            if not phone:
                missing.append('Contact Number')
            if values.get('opening_balance') in (None, ''):
                missing.append('Opening Balance')
            if missing:
                raise UserError(_(
                    'Row %s: mandatory fields missing: %s'
                ) % (row_number, ', '.join(missing)))

            try:
                opening_balance = float(str(values['opening_balance']).replace(',', '').strip())
            except ValueError as exc:
                raise UserError(_(
                    'Row %s: invalid opening balance "%s".'
                ) % (row_number, values['opening_balance'])) from exc

            existing = self.env['ac.customer.master'].search([
                ('contact_number', '=', phone),
                ('company_id', '=', self.company_id.id),
            ], limit=1)

            country = self._resolve_country(values.get('country'))
            state = self._resolve_state(values.get('state'), country)

            lines_vals.append({
                'wizard_id': self.id,
                'row_number': row_number,
                'name': name,
                'contact_number': phone,
                'opening_balance': opening_balance,
                'customer_company': values.get('customer_company') or False,
                'email': values.get('email') or False,
                'aadhar_no': values.get('aadhar_no') or False,
                'pan_no': values.get('pan_no') or False,
                'house_no': values.get('house_no') or False,
                'street': values.get('street') or False,
                'street2': values.get('street2') or False,
                'city': values.get('city') or False,
                'state_id': state.id if state else False,
                'country_id': country.id if country else False,
                'zip': values.get('zip') or False,
                'bank_account_number': values.get('bank_account_number') or False,
                'bank_ifsc_code': values.get('bank_ifsc_code') or False,
                'existing_customer_id': existing.id,
                'action': 'update' if existing else 'create',
            })

        if not lines_vals:
            raise UserError(_('No valid customer rows found in the file.'))

        self.env['ac.customer.import.wizard.line'].create(lines_vals)
        self.state = 'preview'
        return self._reopen()

    def _read_import_rows(self):
        self.ensure_one()
        try:
            content = base64.b64decode(self.data_file)
        except Exception as exc:
            raise UserError(_('Could not read the uploaded file.')) from exc

        filename = (self.filename or '').lower()
        if filename.endswith(('.xlsx', '.xlsm')):
            return self._read_xlsx_rows(content)
        return self._read_csv_rows(content)

    def _read_csv_rows(self, content):
        try:
            text = content.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = content.decode('latin-1')
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise UserError(_('The CSV file is empty or has no header row.'))
        return [dict(row) for row in reader]

    def _read_xlsx_rows(self, content):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise UserError(_(
                'Excel import requires openpyxl. Save the file as CSV, or install openpyxl.'
            )) from exc

        workbook = load_workbook(filename=io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            headers = next(rows_iter)
        except StopIteration as exc:
            raise UserError(_('The Excel file is empty.')) from exc

        headers = [(str(h).strip() if h is not None else '') for h in headers]
        if not any(headers):
            raise UserError(_('The Excel file has no header row.'))

        rows = []
        for values in rows_iter:
            if values is None or all(v is None or str(v).strip() == '' for v in values):
                continue
            row = {}
            for index, header in enumerate(headers):
                if not header:
                    continue
                value = values[index] if index < len(values) else None
                row[header] = '' if value is None else value
            rows.append(row)
        return rows

    @api.model
    def _normalize_header(self, header):
        return (header or '').strip().lower().replace(' ', '_').replace('/', '_')

    @api.model
    def _map_headers(self, headers):
        normalized = {
            self._normalize_header(header): header
            for header in headers
            if header
        }
        mapping = {}
        for field_key, aliases in HEADER_ALIASES.items():
            for alias in aliases:
                if alias in normalized:
                    mapping[field_key] = normalized[alias]
                    break
            else:
                # Also match exact template labels normalized
                for _key, label in IMPORT_COLUMNS:
                    if _key == field_key and self._normalize_header(label) in normalized:
                        mapping[field_key] = normalized[self._normalize_header(label)]
                        break
        return mapping

    @api.model
    def _extract_row_values(self, row, header_map):
        values = {}
        for field_key, header in header_map.items():
            raw = row.get(header)
            if raw is None:
                values[field_key] = ''
            else:
                values[field_key] = str(raw).strip() if not isinstance(raw, (int, float)) else raw
        return values

    def _resolve_country(self, value):
        if not value:
            return self.env['res.country']
        value = str(value).strip()
        Country = self.env['res.country']
        country = Country.search([('code', '=ilike', value)], limit=1)
        if not country:
            country = Country.search([('name', '=ilike', value)], limit=1)
        return country

    def _resolve_state(self, value, country):
        if not value:
            return self.env['res.country.state']
        value = str(value).strip()
        State = self.env['res.country.state']
        domain = ['|', ('name', '=ilike', value), ('code', '=ilike', value)]
        if country:
            domain = [('country_id', '=', country.id)] + domain
        return State.search(domain, limit=1)

    def action_back(self):
        self.ensure_one()
        self.line_ids.unlink()
        self.state = 'upload'
        return self._reopen()

    def action_import(self):
        self.ensure_one()
        if not self.line_ids:
            raise UserError(_('Load a file before importing.'))

        Customer = self.env['ac.customer.master']
        Statement = self.env['ac.customer.statement.line']
        created = 0
        updated = 0
        today = fields.Date.context_today(self)
        currency = self.company_id.currency_id
        india = self.env.ref('base.in', raise_if_not_found=False)

        for line in self.line_ids:
            vals = {
                'name': line.name,
                'contact_number': line.contact_number,
                'company_id': self.company_id.id,
                'customer_company': line.customer_company,
                'email': line.email,
                'aadhar_no': line.aadhar_no,
                'pan_no': line.pan_no,
                'house_no': line.house_no,
                'street': line.street,
                'street2': line.street2,
                'city': line.city,
                'state_id': line.state_id.id,
                'country_id': line.country_id.id or (india.id if india else False),
                'zip': line.zip,
                'bank_account_number': line.bank_account_number,
                'bank_ifsc_code': line.bank_ifsc_code,
            }
            customer = line.existing_customer_id
            if customer:
                customer.write(vals)
                updated += 1
            else:
                customer = Customer.create(vals)
                created += 1

            # Re-import replaces previous imported opening balances for this customer.
            Statement.search([
                ('customer_id', '=', customer.id),
                ('company_id', '=', self.company_id.id),
                ('transaction_type', '=', 'opening'),
                ('state', '=', 'posted'),
            ]).write({'state': 'cancelled'})

            amount = abs(float_round(line.opening_balance, precision_rounding=currency.rounding))
            if float_is_zero(amount, precision_rounding=currency.rounding):
                # Zero opening balance still cancels old openings; no new line.
                continue

            debit = amount if line.opening_balance > 0 else 0.0
            credit = amount if line.opening_balance < 0 else 0.0
            Statement.create({
                'name': _('Opening Balance'),
                'date': today,
                'customer_id': customer.id,
                'company_id': self.company_id.id,
                'transaction_type': 'opening',
                'debit': debit,
                'credit': credit,
                'notes': _('Imported opening balance'),
            })

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Import Complete'),
                'message': _(
                    'Customers created: %(created)s, updated: %(updated)s.'
                ) % {'created': created, 'updated': updated},
                'type': 'success',
                'sticky': False,
                'next': {'type': 'ir.actions.act_window_close'},
            },
        }

    def _reopen(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }


class AcCustomerImportWizardLine(models.TransientModel):
    _name = 'ac.customer.import.wizard.line'
    _description = 'Customer Import Preview Line'
    _order = 'row_number'

    wizard_id = fields.Many2one(
        'ac.customer.import.wizard',
        required=True,
        ondelete='cascade',
    )
    row_number = fields.Integer(string='Row')
    name = fields.Char(string='Customer Name', required=True)
    contact_number = fields.Char(string='Contact Number', required=True)
    opening_balance = fields.Float(string='Opening Balance', required=True)
    customer_company = fields.Char(string='Customer Company')
    email = fields.Char(string='Email')
    aadhar_no = fields.Char(string='Aadhar Card No')
    pan_no = fields.Char(string='PAN Card No')
    house_no = fields.Char(string='House No/Flat No')
    street = fields.Char(string='Street Name')
    street2 = fields.Char(string='Street 2')
    city = fields.Char(string='City')
    state_id = fields.Many2one('res.country.state', string='State')
    country_id = fields.Many2one('res.country', string='Country')
    zip = fields.Char(string='PINCODE')
    bank_account_number = fields.Char(string='Account Number')
    bank_ifsc_code = fields.Char(string='IFSC Code')
    existing_customer_id = fields.Many2one('ac.customer.master', string='Existing Customer')
    action = fields.Selection(
        selection=[
            ('create', 'Create'),
            ('update', 'Update'),
        ],
        string='Action',
        required=True,
    )
